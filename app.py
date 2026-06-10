"""
Extracteur NF Habitat — Interface web
"""

import time, json, threading, os, re
from datetime import datetime
from flask import Flask, render_template_string, jsonify, request, send_file, Response
import pandas as pd
import requests
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
import io

app = Flask(__name__)

state = {
    "status": "idle",
    "message": "",
    "progress": 0,
    "data": [],
    "total": 0,
}

QLIK_APP = "ed735054-1aec-4957-ad1b-c531be3a90bd"
QLIK_URI = "https://qlik-public.nf-habitat.fr"
URL      = "https://www.nf-habitat.fr/moteur-de-recherche-des-operations-certifiees-nf-habitat/"
CERQUAL_PDF = "https://api.cerqual-pro.net/v1/qualitel_site_service/certificats/{ref}"


def extract_year(value):
    """Extrait l'annee (4 chiffres) d'une date au format texte (ex: '15/03/2022' -> '2022')."""
    if not value:
        return ""
    m = re.search(r"(?:19|20)\d{2}", str(value))
    return m.group(0) if m else ""

SETUP_JS = """
window.__qlikReady = false; window.__qlikError = null;
(async () => {
  try {
    const r = await fetch('""" + QLIK_URI + """/qps/csrftoken', {credentials:'include'});
    const csrf = r.headers.get('qlik-csrf-token') || '';
    const schema = await (await fetch('https://unpkg.com/enigma.js@2.14.0/schemas/12.2015.0.json')).json();
    const wsUrl = '""" + QLIK_URI.replace('https','wss') + """/app/""" + QLIK_APP + """' + (csrf ? '?qlik-csrf-token='+csrf : '');
    const session = window.enigma.create({url: wsUrl, schema: schema});
    const global = await session.open();
    window.__qlikApp = await global.openDoc('""" + QLIK_APP + """');
    window.__qlikReady = true;
  } catch(e) { window.__qlikError = e.toString(); }
})();
"""

def build_js(filters):
    fields = [
        "Num\u00e9ro de contrat", "Nom commercial", "Code postal", "Ville",
        "D\u00e9partement", "Statut de certification", "Promoteur Cabinet",
        "Promoteur Groupe", "URL Promoteur", "Date enregistrement"
    ]
    parts = []
    if filters.get("marque"):
        parts.append("[Marque]={'" + filters["marque"] + "'}")
    if filters.get("type"):
        parts.append("[Type de logement]={'" + filters["type"] + "'}")
    if filters.get("region"):
        parts.append("[R\u00e9gion]={'" + filters["region"] + "'}")
    statut = filters.get("statut", "both")
    if statut == "both":
        parts.append('[Statut de certification]={"Certifi\u00e9e","En cours d\u2019\u00e9valuation"}')
    elif statut == "certifiee":
        parts.append('[Statut de certification]={"Certifi\u00e9e"}')
    elif statut == "encours":
        parts.append('[Statut de certification]={"En cours d\u2019\u00e9valuation"}')
    set_expr = "<" + ", ".join(parts) + ">" if parts else "<>"
    fields_json = json.dumps(fields, ensure_ascii=False)
    set_json    = json.dumps(set_expr, ensure_ascii=False)
    return """
const cb = arguments[arguments.length - 1];
(async () => {
  try {
    const app = window.__qlikApp;
    const fields = """ + fields_json + """;
    const setExpr = """ + set_json + """;
    const dims = fields.map(f => ({qDef: {qFieldDefs: [f]}, qNullSuppression: false}));
    const obj = await app.createSessionObject({
      qInfo: {qType: 'hypercube'},
      qHyperCubeDef: {
        qInitialDataFetch: [{qHeight: 900, qWidth: 11}],
        qDimensions: dims,
        qMeasures: [{qDef: {qDef: "sum({" + setExpr + "} 1)"}, qLabel: "s"}],
        qSuppressZero: true, qSuppressMissing: false, qMode: 'S', qStateName: '$'
      }
    });
    window.__listObj = obj;
    const layout = await obj.getLayout();
    const size = layout.qHyperCube.qSize;
    cb({
      totalRows: size.qcy,
      firstPage: layout.qHyperCube.qDataPages[0].qMatrix.map(r => r.map(c => c.qText))
    });
  } catch(e) { cb({error: e.toString()}); }
})();
"""

PAGE_JS = """
const top = arguments[0], h = arguments[1], cb = arguments[arguments.length-1];
(async () => {
  try {
    const pages = await window.__listObj.getHyperCubeData('/qHyperCubeDef',
      [{qTop: top, qLeft: 0, qHeight: h, qWidth: 11}]);
    cb(pages[0].qMatrix.map(r => r.map(c => c.qText)));
  } catch(e) { cb({error: e.toString()}); }
})();
"""

def get_driver():
    opts = Options()
    opts.add_argument("--headless=new")
    opts.add_argument("--window-size=1400,900")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    if os.path.exists("/root/.nix-profile/bin/chromium"):
        opts.binary_location = "/root/.nix-profile/bin/chromium"
        return webdriver.Chrome(service=Service("/root/.nix-profile/bin/chromedriver"), options=opts)
    else:
        from webdriver_manager.chrome import ChromeDriverManager
        return webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=opts)

def run_extraction(filters):
    state.update({"status": "running", "message": "Ouverture du navigateur...", "progress": 5, "data": [], "total": 0})
    driver = None
    try:
        driver = get_driver()
        driver.set_script_timeout(90)

        state.update({"message": "Connexion au site NF Habitat...", "progress": 15})
        driver.get(URL)
        time.sleep(6)
        try:
            driver.execute_script("document.querySelectorAll('[class*=axeptio],[id*=axeptio]').forEach(e=>e.remove());")
        except: pass

        state.update({"message": "Connexion \u00e0 la base Cerqual...", "progress": 25})
        driver.execute_script(SETUP_JS)

        for i in range(40):
            ready = driver.execute_script("return window.__qlikReady;")
            err   = driver.execute_script("return window.__qlikError;")
            if ready: break
            if err:
                state.update({"status": "error", "message": f"Erreur Qlik : {err}"}); return
            time.sleep(1)

        state.update({"message": "R\u00e9cup\u00e9ration des donn\u00e9es...", "progress": 40})
        res = driver.execute_async_script(build_js(filters))
        if isinstance(res, dict) and "error" in res:
            state.update({"status": "error", "message": res["error"]}); return

        total = res["totalRows"]
        state.update({"message": f"{total} op\u00e9rations trouv\u00e9es \u2014 chargement...", "progress": 50, "total": total})

        matrix = res["firstPage"]
        fetched = len(matrix)
        while fetched < total:
            page = driver.execute_async_script(PAGE_JS, fetched, min(900, total - fetched))
            if isinstance(page, dict) and "error" in page: break
            matrix.extend(page)
            fetched = len(matrix)
            pct = 50 + int((fetched / total) * 40)
            state.update({"message": f"{fetched}/{total} r\u00e9cup\u00e9r\u00e9es...", "progress": pct})

        seen = set()
        records = []
        for r in matrix:
            num = r[0] if r else ""
            if num in seen: continue
            seen.add(num)
            statut = r[5] if len(r) > 5 else ""
            date_val = r[9] if len(r) > 9 else ""
            # Lien PDF uniquement pour les certifiées
            pdf_url = CERQUAL_PDF.format(ref=num) if statut == "Certifi\u00e9e" and num else ""
            records.append({
                "reference":     num,
                "nom":           r[1] if len(r) > 1 else "",
                "cp":            r[2] if len(r) > 2 else "",
                "ville":         r[3] if len(r) > 3 else "",
                "departement":   r[4] if len(r) > 4 else "",
                "statut":        statut,
                "promoteur":     r[6] if (len(r) > 6 and r[6] not in ("", "-")) else (r[7] if len(r) > 7 else ""),
                "groupe":        r[7] if len(r) > 7 else "",
                "url_promoteur": r[8] if len(r) > 8 else "",
                "date":          date_val,
                "year":          extract_year(date_val),
                "pdf":           pdf_url,
            })

        state.update({"status": "done", "message": f"\u2705 {len(records)} op\u00e9rations extraites", "progress": 100, "data": records})

    except Exception as e:
        state.update({"status": "error", "message": str(e)})
    finally:
        if driver:
            try: driver.quit()
            except: pass

@app.route("/")
def index():
    return render_template_string(HTML)

@app.route("/api/extract", methods=["POST"])
def extract():
    if state["status"] == "running":
        return jsonify({"error": "Extraction d\u00e9j\u00e0 en cours"}), 400
    filters = request.json or {}
    t = threading.Thread(target=run_extraction, args=(filters,))
    t.daemon = True
    t.start()
    return jsonify({"ok": True})

@app.route("/api/status")
def status():
    return jsonify({
        "status":   state["status"],
        "message":  state["message"],
        "progress": state["progress"],
        "total":    state["total"],
        "count":    len(state["data"]),
        "preview":  state["data"][:20],
    })

@app.route("/api/pdf/<reference>")
def pdf_proxy(reference):
    # Securite : on n'accepte que des references alphanumeriques (pas d'injection de chemin)
    if not reference or not re.fullmatch(r"[A-Za-z0-9_-]+", reference):
        return "R\u00e9f\u00e9rence invalide", 400
    url = CERQUAL_PDF.format(ref=reference)
    try:
        r = requests.get(url, timeout=30, headers={"User-Agent": "Mozilla/5.0"})
    except Exception as e:
        return f"Erreur de connexion au certificat : {e}", 502
    if r.status_code != 200 or not r.content:
        return f"Certificat introuvable ({r.status_code})", 404
    # Content-Type pdf + Content-Disposition inline => affichage direct dans l'onglet
    return Response(
        r.content,
        mimetype="application/pdf",
        headers={
            "Content-Disposition": f'inline; filename="{reference}.pdf"',
            "Cache-Control": "public, max-age=3600",
        },
    )


@app.route("/api/download")
def download():
    if not state["data"]: return "Aucune donn\u00e9e", 400

    rows = state["data"]

    # --- Filtre annee : on garde les operations dont l'annee est >= annee choisie ---
    year_filter = (request.args.get("year") or "").strip()
    if year_filter.isdigit():
        ymin = int(year_filter)
        rows = [r for r in rows if (r.get("year") or "").isdigit() and int(r["year"]) >= ymin]

    if not rows:
        return "Aucune donn\u00e9e pour ce filtre", 400

    # Base URL publique (pour que les liens PDF de l'Excel pointent vers le proxy en ligne)
    host = request.host
    scheme = "http" if ("localhost" in host or "127.0.0.1" in host) else "https"
    base = f"{scheme}://{host}/"

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Op\u00e9rations"

    headers = ["R\u00e9f\u00e9rence", "Nom op\u00e9ration", "Code postal", "Ville", "D\u00e9partement",
               "Statut", "Promoteur", "Groupe promoteur", "Site web promoteur",
               "Date enregistrement", "Ann\u00e9e", "Certificat PDF"]
    ws.append(headers)

    # Palette / styles
    header_fill = PatternFill("solid", fgColor="2E7D32")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center      = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center_s    = Alignment(horizontal="center", vertical="center")
    left        = Alignment(horizontal="left", vertical="center")
    thin        = Side(style="thin", color="D9E4D9")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_alt    = PatternFill("solid", fgColor="F1F8E9")
    fill_certif = PatternFill("solid", fgColor="C8E6C9")
    fill_cours  = PatternFill("solid", fgColor="FFF3C4")
    link_font   = Font(color="1565C0", underline="single")

    STATUS_COL, YEAR_COL, PDF_COL = 6, 11, 12

    # Ligne d'en-tete
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border
    ws.row_dimensions[1].height = 30

    # Lignes de donnees
    for i, rec in enumerate(rows, start=2):
        values = [
            rec.get("reference", ""), rec.get("nom", ""), rec.get("cp", ""),
            rec.get("ville", ""), rec.get("departement", ""), rec.get("statut", ""),
            rec.get("promoteur", ""), rec.get("groupe", ""), rec.get("url_promoteur", ""),
            rec.get("date", ""), rec.get("year", ""), "",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = border
            cell.alignment = center_s if c in (1, 3, 5, 11) else left
            if i % 2 == 0:
                cell.fill = fill_alt

        # Statut colore
        st = ws.cell(row=i, column=STATUS_COL)
        if rec.get("statut") == "Certifi\u00e9e":
            st.fill = fill_certif
        elif rec.get("statut"):
            st.fill = fill_cours
        st.alignment = center_s

        # Lien PDF -> proxy /api/pdf/<ref> (ouverture inline dans le navigateur)
        pdf_cell = ws.cell(row=i, column=PDF_COL)
        if rec.get("pdf") and rec.get("reference"):
            pdf_cell.value = "Voir le certificat"
            pdf_cell.hyperlink = f"{base}api/pdf/{rec['reference']}"
            pdf_cell.font = link_font
        else:
            pdf_cell.value = "\u2014"
        pdf_cell.alignment = center_s

    # Largeurs de colonnes
    widths = [16, 36, 11, 18, 16, 14, 26, 24, 30, 18, 8, 20]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # Filtre auto + gel de la ligne d'en-tete
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    suffix = f"_des_{year_filter}" if year_filter.isdigit() else ""
    return send_file(buf, as_attachment=True,
                     download_name=f"NF_Habitat{suffix}_{stamp}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

HTML = '''<!DOCTYPE html>
<html lang="fr"><head><meta charset="UTF-8"/><meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>Extracteur NF Habitat</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:'Segoe UI',Arial,sans-serif;background:#f0f4f0;color:#1a1a1a;min-height:100vh}
header{background:#2E7D32;color:white;padding:18px 32px;display:flex;align-items:center;gap:14px;box-shadow:0 2px 8px rgba(0,0,0,.15)}
header h1{font-size:1.3rem;font-weight:700}header p{font-size:.82rem;opacity:.8;margin-top:2px}
.wrap{max-width:1100px;margin:0 auto;padding:24px 16px}
.card{background:white;border-radius:10px;box-shadow:0 1px 6px rgba(0,0,0,.08);padding:24px 28px;margin-bottom:20px}
.card h2{font-size:.9rem;font-weight:700;color:#2E7D32;text-transform:uppercase;letter-spacing:.05em;margin-bottom:16px}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(200px,1fr));gap:14px}
label{display:block;font-size:.75rem;font-weight:600;color:#666;text-transform:uppercase;letter-spacing:.04em;margin-bottom:4px}
select,input{width:100%;padding:9px 12px;border:1.5px solid #ddd;border-radius:6px;font-size:.9rem;background:#f9f9f9}
select:focus,input:focus{outline:none;border-color:#4CAF50;background:white}
.actions{display:flex;gap:12px;flex-wrap:wrap;margin-top:18px;align-items:center}
.btn{padding:10px 22px;border-radius:7px;border:none;font-size:.9rem;font-weight:600;cursor:pointer;display:flex;align-items:center;gap:7px;transition:all .15s}
.btn-green{background:#2E7D32;color:white}.btn-green:hover{background:#1B5E20}.btn-green:disabled{background:#A5D6A7;cursor:not-allowed}
.btn-blue{background:#1565C0;color:white}.btn-blue:hover{background:#0D47A1}.btn-blue:disabled{background:#90CAF9;cursor:not-allowed}
.btn-gray{background:#e0e0e0;color:#333}.btn-gray:hover{background:#bdbdbd}
#status-box{padding:10px 16px;border-radius:7px;font-size:.88rem;font-weight:500;display:none;margin-top:12px}
.s-loading{background:#FFF9C4;color:#E65100;display:block!important}
.s-done{background:#E8F5E9;color:#2E7D32;display:block!important}
.s-error{background:#FFEBEE;color:#C62828;display:block!important}
.progress-wrap{height:6px;background:#e0e0e0;border-radius:4px;margin-top:10px;display:none}
.progress-bar{height:6px;background:#4CAF50;border-radius:4px;width:0%;transition:width .3s}
.stats{display:flex;gap:20px;flex-wrap:wrap;padding:10px 0;font-size:.85rem;color:#555}
.stat strong{color:#2E7D32;font-size:1.1rem}
.table-wrap{overflow-x:auto;overflow-y:auto;max-height:500px;margin-top:8px}
table{width:100%;border-collapse:collapse;font-size:.83rem}
thead th{background:#2E7D32;color:white;padding:9px 11px;text-align:left;white-space:nowrap;position:sticky;top:0;z-index:1}
tbody tr:nth-child(even){background:#F1F8E9}tbody tr:hover{background:#DCEDC8}
tbody td{padding:7px 11px;border-bottom:1px solid #e0e0e0;vertical-align:middle}
.badge{display:inline-block;padding:2px 8px;border-radius:10px;font-size:.75rem;font-weight:600}
.b-c{background:#C8E6C9;color:#1B5E20}.b-e{background:#FFF9C4;color:#E65100}
a.pdf-link{color:#1565C0;font-size:.8rem;text-decoration:none;white-space:nowrap}
a.pdf-link:hover{text-decoration:underline}
tfoot td{padding:10px 11px;color:#888;font-size:.8rem;font-style:italic}
</style></head><body>
<header>
  <svg width="34" height="34" viewBox="0 0 34 34" fill="none">
    <rect width="34" height="34" rx="7" fill="white" fill-opacity=".15"/>
    <path d="M17 5L29 12V27H5V12Z" stroke="white" stroke-width="2" fill="none"/>
    <rect x="13" y="19" width="8" height="8" fill="white" fill-opacity=".8"/>
    <circle cx="17" cy="14" r="3" fill="white"/>
  </svg>
  <div><h1>Extracteur NF Habitat</h1><p>Donn\u00e9es live Cerqual \u2014 certifi\u00e9es et en cours d\u2019\u00e9valuation</p></div>
</header>
<div class="wrap">
  <div class="card">
    <h2>\U0001f50d Filtres</h2>
    <div class="grid">
      <div><label>R\u00e9gion</label><select id="f-region">
        <option value="">Toutes</option>
        <option value="Ile-de-France" selected>\u00cele-de-France</option>
        <option value="Auvergne-Rh\u00f4ne-Alpes">Auvergne-Rh\u00f4ne-Alpes</option>
        <option value="Bretagne">Bretagne</option>
        <option value="Grand Est">Grand Est</option>
        <option value="Hauts-de-France">Hauts-de-France</option>
        <option value="Normandie">Normandie</option>
        <option value="Nouvelle-Aquitaine">Nouvelle-Aquitaine</option>
        <option value="Occitanie">Occitanie</option>
        <option value="Pays de la Loire">Pays de la Loire</option>
        <option value="Provence-Alpes-C\u00f4te d\u2019Azur">Provence-Alpes-C\u00f4te d\u2019Azur</option>
      </select></div>
      <div><label>Type de logement</label><select id="f-type">
        <option value="">Tous</option>
        <option value="Collectif" selected>Collectif</option>
        <option value="Individuel">Individuel</option>
      </select></div>
      <div><label>Certification</label><select id="f-marque">
        <option value="">Toutes</option>
        <option value="NF Habitat HQE" selected>NF Habitat HQE</option>
        <option value="NF Habitat">NF Habitat</option>
      </select></div>
      <div><label>Statut</label><select id="f-statut">
        <option value="both" selected>Certifi\u00e9e + En cours</option>
        <option value="certifiee">Certifi\u00e9e uniquement</option>
        <option value="encours">En cours uniquement</option>
      </select></div>
      <div><label>Ann\u00e9e (\u00e0 partir de)</label><select id="f-year">
        <option value="" selected>Toutes les ann\u00e9es</option>
        <option value="2027">2027</option>
        <option value="2026">2026</option>
        <option value="2025">2025</option>
        <option value="2024">2024</option>
        <option value="2023">2023</option>
        <option value="2022">2022</option>
        <option value="2021">2021</option>
        <option value="2020">2020</option>
        <option value="2019">2019</option>
        <option value="2018">2018</option>
        <option value="2017">2017</option>
        <option value="2016">2016</option>
        <option value="2015">2015</option>
        <option value="2014">2014</option>
        <option value="2013">2013</option>
        <option value="2012">2012</option>
      </select></div>
      <div><label>Recherche libre</label><input type="text" id="f-search" placeholder="nom, ville, promoteur..."/></div>
    </div>
    <div class="actions">
      <button class="btn btn-green" id="btn-go" onclick="lancer()">&#9654; Lancer l\u2019extraction</button>
      <button class="btn btn-gray" onclick="reset()">\u2715 R\u00e9initialiser</button>
      <button class="btn btn-blue" id="btn-dl" disabled onclick="telecharger()">\u2b07 T\u00e9l\u00e9charger Excel</button>
    </div>
    <div class="progress-wrap" id="pw"><div class="progress-bar" id="pb"></div></div>
    <div id="status-box"></div>
  </div>
  <div class="card" id="results-card" style="display:none">
    <h2>&#128203; R\u00e9sultats <span id="subtitle" style="font-weight:400;text-transform:none;font-size:.85rem;color:#888"></span></h2>
    <div class="stats" id="stats"></div>
    <div class="table-wrap"><table>
      <thead><tr>
        <th>R\u00e9f\u00e9rence</th><th>Nom op\u00e9ration</th><th>Promoteur</th>
        <th>Statut</th><th>CP</th><th>Ville</th><th>D\u00e9partement</th><th>Date enreg.</th><th>Certificat</th>
      </tr></thead>
      <tbody id="tbody"></tbody>
      <tfoot><tr><td colspan="9" id="tfoot-msg"></td></tr></tfoot>
    </table></div>
  </div>
</div>
<script>
let allData=[],timer=null,totalCount=0;
function lancer(){
  const f={region:document.getElementById("f-region").value,type:document.getElementById("f-type").value,
    marque:document.getElementById("f-marque").value,statut:document.getElementById("f-statut").value};
  document.getElementById("btn-go").disabled=true;
  document.getElementById("btn-dl").disabled=true;
  document.getElementById("results-card").style.display="none";
  allData=[];totalCount=0;
  fetch("/api/extract",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(f)})
    .then(()=>{timer=setInterval(pollStatus,1000);});
}
function pollStatus(){
  fetch("/api/status").then(r=>r.json()).then(d=>{
    setStatus(d.status==="running"?"loading":d.status==="done"?"done":"error",d.message);
    setProgress(d.progress);
    if(d.status==="done"||d.status==="error"){
      clearInterval(timer);document.getElementById("btn-go").disabled=false;
      if(d.status==="done"){allData=d.preview;totalCount=d.count;afficher();document.getElementById("btn-dl").disabled=false;}
    }
  });
}
function afficher(){
  const q=document.getElementById("f-search").value.toLowerCase().trim();
  const yr=document.getElementById("f-year").value;
  let data=allData;
  if(yr) data=data.filter(r=>r.year && parseInt(r.year,10)>=parseInt(yr,10));
  if(q) data=data.filter(r=>(r.nom+r.ville+r.promoteur+r.cp).toLowerCase().includes(q));
  const certif=allData.filter(r=>r.statut==="Certifi\u00e9e").length;
  const cours=allData.filter(r=>r.statut&&r.statut.includes("cours")).length;
  document.getElementById("stats").innerHTML=
    `<div class="stat">Total\u00a0: <strong>${totalCount}</strong></div>
     <div class="stat">\u2705 Certifi\u00e9es\u00a0: <strong>${certif}</strong></div>
     <div class="stat">\u23f3 En cours\u00a0: <strong>${cours}</strong></div>`;
  document.getElementById("subtitle").textContent="\u2014 "+totalCount+" op\u00e9rations"+(yr?(" (\u2265 "+yr+")"):"");
  document.getElementById("tbody").innerHTML=data.map(r=>`<tr>
    <td><code style="font-size:.78rem">${r.reference}</code></td>
    <td><strong>${r.nom||"\u2014"}</strong></td>
    <td>${r.promoteur||"\u2014"}</td>
    <td>${r.statut==="Certifi\u00e9e"?'<span class="badge b-c">\u2705 Certifi\u00e9e</span>':'<span class="badge b-e">\u23f3 En cours</span>'}</td>
    <td>${r.cp}</td><td>${r.ville}</td><td>${r.departement}</td><td>${r.year||"\u2014"}</td>
    <td>${r.pdf?`<a class="pdf-link" href="/api/pdf/${r.reference}" target="_blank" rel="noopener">Voir le certificat</a>`:"\u2014"}</td>
  </tr>`).join("");
  document.getElementById("tfoot-msg").textContent = yr
    ? "Filtre ann\u00e9e \u2265 "+yr+" appliqu\u00e9 \u00e0 l\u2019aper\u00e7u \u2014 l\u2019export Excel applique ce filtre sur la totalit\u00e9 des "+totalCount+" op\u00e9rations."
    : (totalCount>20?"Affichage des 20 premi\u00e8res lignes sur "+totalCount+" \u2014 t\u00e9l\u00e9chargez l\u2019Excel pour tout voir.":"");
  document.getElementById("results-card").style.display="block";
}
document.getElementById("f-search").addEventListener("input",()=>{if(allData.length)afficher();});
document.getElementById("f-year").addEventListener("change",()=>{if(allData.length)afficher();});
function telecharger(){
  const y=document.getElementById("f-year").value;
  window.location.href="/api/download"+(y?("?year="+encodeURIComponent(y)):"");
}
function setStatus(type,msg){
  const el=document.getElementById("status-box");
  el.className=type==="loading"?"s-loading":type==="done"?"s-done":"s-error";
  el.innerHTML=msg;el.style.display="block";
}
function setProgress(pct){
  document.getElementById("pw").style.display="block";
  document.getElementById("pb").style.width=pct+"%";
  if(pct>=100)setTimeout(()=>document.getElementById("pw").style.display="none",800);
}
function reset(){
  document.getElementById("f-region").value="Ile-de-France";
  document.getElementById("f-type").value="Collectif";
  document.getElementById("f-marque").value="NF Habitat HQE";
  document.getElementById("f-statut").value="both";
  document.getElementById("f-year").value="";
  document.getElementById("f-search").value="";
  document.getElementById("results-card").style.display="none";
  document.getElementById("status-box").style.display="none";
  document.getElementById("btn-dl").disabled=true;
  allData=[];totalCount=0;
}
</script></body></html>'''

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    is_local = port == 5000
    print("="*50)
    print("  Extracteur NF Habitat")
    print(f"  http://localhost:{port}")
    print("  Ctrl+C pour arreter")
    print("="*50)
    if is_local:
        import webbrowser, threading
        threading.Timer(1.5, lambda: webbrowser.open(f"http://localhost:{port}")).start()
    app.run(host="0.0.0.0", port=port, debug=False)