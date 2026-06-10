"""
Extracteur Prestaterre BEE — Interface web
Appelle directement l'API interne du site (postCertifiedOperations) via Selenium,
exactement comme le fait le site lui-même. Mêmes filtres, même pagination.
"""

import time, re, io, os, threading
from datetime import datetime
from flask import Flask, render_template_string, jsonify, request, send_file

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

state = {
    "status":   "idle",
    "message":  "",
    "progress": 0,
    "data":     [],
    "total":    0,
}

URL_PRESTATERRE = "https://www.prestaterre.eu/operations-certifiees"

# Appelle l'API interne du site pour UNE page et renvoie les lignes (sans lat/long)
PAGE_FETCH_JS = """
const cb = arguments[arguments.length - 1];
const p  = arguments[0];
(async () => {
  try {
    if (typeof postAPI !== 'function' || typeof postCertifiedOperations === 'undefined') {
      cb({error: "API du site indisponible"}); return;
    }
    const data = {
      page: p.page, lastId: p.lastId, navFlag: p.navFlag,
      department: p.department, city: p.city, operationName: p.operationName,
      repositoryAndVersion: p.repositoryAndVersion, mentionsAndLevels: p.mentionsAndLevels
    };
    const resp = await postAPI(postCertifiedOperations, data);
    const rows = (resp && resp.data) ? resp.data.map(function(row){
      const out = [];
      for (const k in row) { if (k !== 'latitude' && k !== 'longitude') out.push(row[k]); }
      return out;
    }) : [];
    cb({
      rows: rows,
      currentPage: resp ? resp.currentPage : p.page,
      totalPages:  resp ? (parseInt(resp.totalPages) || 1) : 1,
      lastId:      resp ? (resp.lastId || "") : ""
    });
  } catch(e) { cb({error: e.toString()}); }
})();
"""


def get_driver():
    opts = Options()
    opts.add_argument("--headless=new")
    opts.add_argument("--window-size=1600,900")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--lang=fr-FR")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    if os.path.exists("/root/.nix-profile/bin/chromium"):
        opts.binary_location = "/root/.nix-profile/bin/chromium"
        return webdriver.Chrome(
            service=Service("/root/.nix-profile/bin/chromedriver"), options=opts)
    from webdriver_manager.chrome import ChromeDriverManager
    return webdriver.Chrome(
        service=Service(ChromeDriverManager().install()), options=opts)


def clean_html(value):
    """Nettoie les valeurs renvoyees par l'API (peuvent contenir <br>, balises...)."""
    if value is None:
        return ""
    s = str(value)
    s = re.sub(r"<\s*br\s*/?\s*>", " / ", s, flags=re.IGNORECASE)
    s = re.sub(r"<[^>]+>", "", s)
    s = re.sub(r"\s*/\s*/\s*", " / ", s)
    return s.strip(" /").strip()


def run_extraction(filters):
    state.update({"status": "running", "message": "Ouverture du navigateur...",
                  "progress": 5, "data": [], "total": 0})
    driver = None
    try:
        driver = get_driver()
        driver.set_script_timeout(60)

        state.update({"message": "Connexion au site Prestaterre...", "progress": 15})
        driver.get(URL_PRESTATERRE)
        time.sleep(4)

        # Fermer un eventuel bandeau cookies (axeptio) — n'empeche pas l'API mais propre
        try:
            driver.execute_script(
                "document.querySelectorAll('[class*=axeptio],[id*=axeptio]').forEach(e=>e.remove());")
        except Exception:
            pass

        # Attendre que les fonctions API du site soient disponibles
        state.update({"message": "Initialisation de l'API du site...", "progress": 25})
        ready = False
        for _ in range(40):
            ready = driver.execute_script(
                "return (typeof postAPI==='function' && typeof postCertifiedOperations!=='undefined');")
            if ready:
                break
            time.sleep(0.5)
        if not ready:
            state.update({"status": "error",
                          "message": "Impossible d'accéder à l'API du site (postAPI non chargé)."})
            return

        # Pagination : on boucle page par page comme le site
        dept = filters.get("department", []) or []
        city = (filters.get("city", "") or "").strip()
        op   = (filters.get("operationName", "") or "").strip()
        ref  = filters.get("repositoryAndVersion", []) or []
        ment = filters.get("mentionsAndLevels", []) or []

        all_rows = []
        page = 1
        last_id = ""
        total_pages = 1
        guard = 0

        while guard < 600:
            guard += 1
            payload = {
                "page": page, "lastId": last_id,
                "navFlag": "next" if page > 1 else "",
                "department": dept, "city": city, "operationName": op,
                "repositoryAndVersion": ref, "mentionsAndLevels": ment,
            }
            res = driver.execute_async_script(PAGE_FETCH_JS, payload)

            if isinstance(res, dict) and res.get("error"):
                if all_rows:
                    break  # on garde ce qu'on a deja
                state.update({"status": "error", "message": res["error"]})
                return

            rows = res.get("rows", [])
            total_pages = res.get("totalPages", 1) or 1
            last_id = res.get("lastId", "") or ""

            for r in rows:
                vals = [clean_html(x) for x in r]
                while len(vals) < 6:
                    vals.append("")
                all_rows.append({
                    "departement":  vals[0],
                    "ville":        vals[1],
                    "operation":    vals[2],
                    "referentiel":  vals[3],
                    "mentions":     vals[4],
                    "fin_validite": vals[5],
                })

            pct = 30 + int((page / max(total_pages, 1)) * 65)
            state.update({
                "message":  f"Page {page}/{total_pages} — {len(all_rows)} opérations...",
                "progress": min(pct, 95),
                "total":    len(all_rows),
            })

            if page >= total_pages or not rows:
                break
            page += 1

        # Deduplication (departement, ville, operation, referentiel)
        seen, deduped = set(), []
        for r in all_rows:
            key = (r["departement"], r["ville"], r["operation"], r["referentiel"])
            if key not in seen:
                seen.add(key)
                deduped.append(r)

        state.update({
            "status": "done",
            "message": f"\u2705 {len(deduped)} opérations extraites",
            "progress": 100,
            "data": deduped,
            "total": len(deduped),
        })

    except Exception as e:
        state.update({"status": "error", "message": str(e)})
    finally:
        if driver:
            try:
                driver.quit()
            except Exception:
                pass


@app.route("/")
def index():
    return render_template_string(HTML)


@app.route("/api/extract", methods=["POST"])
def extract():
    if state["status"] == "running":
        return jsonify({"error": "Extraction déjà en cours"}), 400
    filters = request.json or {}
    t = threading.Thread(target=run_extraction, args=(filters,))
    t.daemon = True
    t.start()
    return jsonify({"ok": True})


@app.route("/api/status")
def status():
    return jsonify({
        "status":   state["status"],
        "message":  state["message"],
        "progress": state["progress"],
        "total":    state["total"],
        "count":    len(state["data"]),
        "preview":  state["data"][:20],
    })


@app.route("/api/download")
def download():
    if not state["data"]:
        return "Aucune donnée", 400
    rows = state["data"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Opérations BEE"

    headers = ["Département", "Ville", "Opération",
               "Référentiel et Version", "Mentions et Niveaux", "Fin de validité"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1B5E20")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin   = Side(style="thin", color="C8E6C9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_alt = PatternFill("solid", fgColor="F9FBE7")

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border
    ws.row_dimensions[1].height = 28

    for i, rec in enumerate(rows, start=2):
        values = [rec.get("departement", ""), rec.get("ville", ""), rec.get("operation", ""),
                  rec.get("referentiel", ""), rec.get("mentions", ""), rec.get("fin_validite", "")]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = border
            cell.alignment = center if c in (1, 6) else left
            if i % 2 == 0:
                cell.fill = fill_alt

    widths = [18, 22, 44, 36, 40, 16]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return send_file(buf, as_attachment=True,
                     download_name=f"Prestaterre_BEE_{stamp}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ───────────────────────── Listes (identiques au site) ─────────────────────────
REFERENTIELS = ["BEE Logement Neuf", "BEE Logement Rénovation", "BEE Tertiaire Neuf",
                "BEE Tertiaire Rénovation", "BEE Tertiaire Exploitation"]
MENTIONS = ["Bâtiment Performance Énergétique", "BEE+", "Habitat qualité", "E+C-",
            "Effinergie", "BBCA", "RT 2012", "HPE", "RE2020"]
DEPARTEMENTS = ["01 - Ain","02 - Aisne","05 - Hautes-Alpes","06 - Alpes-Maritimes","07 - Ardèche",
    "08 - Ardennes","10 - Aube","12 - Aveyron","13 - Bouches-du-Rhône","14 - Calvados","16 - Charente",
    "17 - Charente-Maritime","18 - Cher","20 - Corse","21 - Côte-d'Or","22 - Côtes-d'Armor","24 - Dordogne",
    "25 - Doubs","26 - Drôme","27 - Eure","28 - Eure-et-Loir","29 - Finistère","2A - Corse-du-Sud",
    "2B - Haute-Corse","30 - Gard","31 - Haute-Garonne","33 - Gironde","34 - Hérault","35 - Ile-et-Vilaine",
    "36 - Indre","37 - Indre-et-Loire","38 - Isère","39 - Jura","40 - Landes","41 - Loir-et-Cher","42 - Loire",
    "44 - Loire-Atlantique","45 - Loiret","47 - Lot-et-Garonne","49 - Maine-et-Loire","50 - Manche","51 - Marne",
    "52 - Haute-Marne","54 - Meurthe-et-Moselle","56 - Morbihan","57 - Moselle","59 - Nord","60 - Oise",
    "62 - Pas-de-Calais","63 - Puy-de-Dôme","64 - Pyrénées-Atlantiques","67 - Bas-Rhin","68 - Haut-Rhin",
    "69 - Rhône","70 - Haute-Saône","71 - Saône-et-Loire","72 - Sarthe","73 - Savoie","74 - Haute-Savoie",
    "75 - Paris","76 - Seine-Maritime","77 - Seine-et-Marne","78 - Yvelines","79 - Deux-Sèvres","80 - Somme",
    "81 - Tarn","83 - Var","84 - Vaucluse","85 - Vendée","86 - Vienne","88 - Vosges","89 - Yonne","91 - Essonne",
    "92 - Hauts-de-Seine","93 - Seine-Saint-Denis","94 - Val-de-Marne","95 - Val-d'Oise"]


def _checkboxes(group_id, items):
    return "".join(
        f'<label class="oc-ms-opt"><input type="checkbox" value="{v}" '
        f'onchange="msLabel(\'{group_id}\')"/> {v}</label>' for v in items)


HTML = '''<!DOCTYPE html>
<html lang="fr"><head>
<meta charset="UTF-8"/><meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>Extracteur Prestaterre BEE</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:'Segoe UI',Arial,sans-serif;background:#f0f4f0;color:#1a1a1a;min-height:100vh}
header{background:#1B5E20;color:white;padding:18px 32px;display:flex;align-items:center;gap:14px;box-shadow:0 2px 8px rgba(0,0,0,.2)}
header h1{font-size:1.3rem;font-weight:700}header p{font-size:.82rem;opacity:.8;margin-top:2px}
.wrap{max-width:1200px;margin:0 auto;padding:24px 16px}
.card{background:white;border-radius:10px;box-shadow:0 1px 6px rgba(0,0,0,.08);padding:24px 28px;margin-bottom:20px}
.card h2{font-size:.9rem;font-weight:700;color:#1B5E20;text-transform:uppercase;letter-spacing:.05em;margin-bottom:16px}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(220px,1fr));gap:14px}
label.fld{display:block;font-size:.75rem;font-weight:600;color:#666;text-transform:uppercase;letter-spacing:.04em;margin-bottom:4px}
input.txt{width:100%;padding:9px 12px;border:1.5px solid #ddd;border-radius:6px;font-size:.9rem;background:#f9f9f9}
input.txt:focus{outline:none;border-color:#4CAF50;background:white}
/* multi-select */
.oc-ms{position:relative}
.oc-ms-btn{width:100%;padding:9px 12px;border:1.5px solid #ddd;border-radius:6px;font-size:.9rem;background:#f9f9f9;cursor:pointer;display:flex;justify-content:space-between;align-items:center;gap:8px;user-select:none}
.oc-ms-btn:hover{border-color:#4CAF50}
.oc-ms-label{white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:#333}
.oc-ms.open .oc-ms-pop{display:block}
.oc-ms-pop{display:none;position:absolute;z-index:30;top:calc(100% + 4px);left:0;right:0;background:white;border:1px solid #cdd;border-radius:6px;box-shadow:0 6px 18px rgba(0,0,0,.15);max-height:260px;overflow-y:auto;padding:6px}
.oc-ms-opt{display:flex;align-items:center;gap:8px;padding:6px 8px;font-size:.85rem;border-radius:4px;cursor:pointer}
.oc-ms-opt:hover{background:#F1F8E9}
.oc-ms-opt input{width:16px;height:16px;cursor:pointer}
.actions{display:flex;gap:12px;flex-wrap:wrap;margin-top:18px;align-items:center}
.btn{padding:10px 22px;border-radius:7px;border:none;font-size:.9rem;font-weight:600;cursor:pointer;display:flex;align-items:center;gap:7px;transition:all .15s}
.btn-green{background:#1B5E20;color:white}.btn-green:hover{background:#2E7D32}.btn-green:disabled{background:#A5D6A7;cursor:not-allowed}
.btn-blue{background:#1565C0;color:white}.btn-blue:hover{background:#0D47A1}.btn-blue:disabled{background:#90CAF9;cursor:not-allowed}
.btn-gray{background:#e0e0e0;color:#333}.btn-gray:hover{background:#bdbdbd}
#status-box{padding:10px 16px;border-radius:7px;font-size:.88rem;font-weight:500;display:none;margin-top:12px}
.s-loading{background:#FFF9C4;color:#E65100;display:block!important}
.s-done{background:#E8F5E9;color:#1B5E20;display:block!important}
.s-error{background:#FFEBEE;color:#C62828;display:block!important}
.progress-wrap{height:6px;background:#e0e0e0;border-radius:4px;margin-top:10px;display:none}
.progress-bar{height:6px;background:#4CAF50;border-radius:4px;width:0%;transition:width .3s}
.stats{display:flex;gap:20px;flex-wrap:wrap;padding:10px 0;font-size:.85rem;color:#555}
.stat strong{color:#1B5E20;font-size:1.1rem}
.table-wrap{overflow-x:auto;overflow-y:auto;max-height:520px;margin-top:8px}
table{width:100%;border-collapse:collapse;font-size:.83rem}
thead th{background:#1B5E20;color:white;padding:9px 11px;text-align:left;white-space:nowrap;position:sticky;top:0;z-index:1}
tbody tr:nth-child(even){background:#F9FBE7}tbody tr:hover{background:#DCEDC8}
tbody td{padding:7px 11px;border-bottom:1px solid #e0e0e0;vertical-align:top}
.badge-ref{display:inline-block;padding:2px 8px;border-radius:10px;font-size:.74rem;font-weight:600;background:#E8F5E9;color:#1B5E20}
tfoot td{padding:10px 11px;color:#888;font-size:.8rem;font-style:italic}
</style></head><body>
<header>
  <svg width="34" height="34" viewBox="0 0 34 34" fill="none">
    <rect width="34" height="34" rx="7" fill="white" fill-opacity=".15"/>
    <path d="M17 4L30 11V25L17 30L4 25V11Z" stroke="white" stroke-width="2" fill="none"/>
    <circle cx="17" cy="17" r="4" fill="white" fill-opacity=".8"/>
  </svg>
  <div><h1>Extracteur Prestaterre BEE</h1>
  <p>Opérations certifiées — données live prestaterre.eu</p></div>
</header>
<div class="wrap">
  <div class="card">
    <h2>🔍 Filtres</h2>
    <div class="grid">
      <div>
        <label class="fld">Nom de l'opération</label>
        <input type="text" class="txt" id="f-op" placeholder="NOM DE L'OPÉRATION"/>
      </div>
      <div>
        <label class="fld">Référentiel et version</label>
        <div class="oc-ms" id="ms-ref">
          <div class="oc-ms-btn" onclick="msToggle('ms-ref')">
            <span class="oc-ms-label" data-ph="REFERENTIEL ET VERSION">REFERENTIEL ET VERSION</span><span>▾</span>
          </div>
          <div class="oc-ms-pop">__REF__</div>
        </div>
      </div>
      <div>
        <label class="fld">Mentions et niveaux</label>
        <div class="oc-ms" id="ms-ment">
          <div class="oc-ms-btn" onclick="msToggle('ms-ment')">
            <span class="oc-ms-label" data-ph="MENTIONS ET NIVEAUX">MENTIONS ET NIVEAUX</span><span>▾</span>
          </div>
          <div class="oc-ms-pop">__MENT__</div>
        </div>
      </div>
      <div>
        <label class="fld">Département</label>
        <div class="oc-ms" id="ms-dept">
          <div class="oc-ms-btn" onclick="msToggle('ms-dept')">
            <span class="oc-ms-label" data-ph="DÉPARTEMENT">DÉPARTEMENT</span><span>▾</span>
          </div>
          <div class="oc-ms-pop">__DEPT__</div>
        </div>
      </div>
      <div>
        <label class="fld">Ville</label>
        <input type="text" class="txt" id="f-ville" placeholder="VILLE"/>
      </div>
    </div>
    <div class="actions">
      <button class="btn btn-green" id="btn-go" onclick="lancer()">&#9654; Lancer l'extraction</button>
      <button class="btn btn-gray" onclick="reset()">✕ Réinitialiser</button>
      <button class="btn btn-blue" id="btn-dl" disabled onclick="telecharger()">⬇ Télécharger Excel</button>
    </div>
    <div class="progress-wrap" id="pw"><div class="progress-bar" id="pb"></div></div>
    <div id="status-box"></div>
  </div>

  <div class="card" id="results-card" style="display:none">
    <h2>📋 Résultats <span id="subtitle" style="font-weight:400;text-transform:none;font-size:.85rem;color:#888"></span></h2>
    <div class="stats" id="stats"></div>
    <div class="table-wrap"><table>
      <thead><tr>
        <th>Département</th><th>Ville</th><th>Opération</th>
        <th>Référentiel et Version</th><th>Mentions et Niveaux</th><th>Fin de validité</th>
      </tr></thead>
      <tbody id="tbody"></tbody>
      <tfoot><tr><td colspan="6" id="tfoot-msg"></td></tr></tfoot>
    </table></div>
  </div>
</div>

<script>
let allData=[], timer=null, totalCount=0;

function msToggle(id){
  const el=document.getElementById(id);
  document.querySelectorAll('.oc-ms.open').forEach(x=>{ if(x.id!==id) x.classList.remove('open'); });
  el.classList.toggle('open');
}
document.addEventListener('click',e=>{
  if(!e.target.closest('.oc-ms')) document.querySelectorAll('.oc-ms.open').forEach(x=>x.classList.remove('open'));
});
function msValues(id){
  return [...document.querySelectorAll('#'+id+' .oc-ms-pop input:checked')].map(c=>c.value);
}
function msLabel(id){
  const v=msValues(id);
  const lbl=document.querySelector('#'+id+' .oc-ms-label');
  lbl.textContent = v.length ? (v.length+' sélectionné'+(v.length>1?'s':'')) : lbl.getAttribute('data-ph');
}

function lancer(){
  const f={
    operationName:        document.getElementById("f-op").value.trim(),
    city:                 document.getElementById("f-ville").value.trim(),
    repositoryAndVersion: msValues('ms-ref'),
    mentionsAndLevels:    msValues('ms-ment'),
    department:           msValues('ms-dept'),
  };
  document.getElementById("btn-go").disabled=true;
  document.getElementById("btn-dl").disabled=true;
  document.getElementById("results-card").style.display="none";
  allData=[]; totalCount=0;
  fetch("/api/extract",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(f)})
    .then(()=>{ timer=setInterval(pollStatus,1500); });
}
function pollStatus(){
  fetch("/api/status").then(r=>r.json()).then(d=>{
    setStatus(d.status==="running"?"loading":d.status==="done"?"done":"error", d.message);
    setProgress(d.progress);
    if(d.status==="done"||d.status==="error"){
      clearInterval(timer);
      document.getElementById("btn-go").disabled=false;
      if(d.status==="done"){ allData=d.preview; totalCount=d.count; afficher();
        document.getElementById("btn-dl").disabled=false; }
    }
  });
}
function afficher(){
  document.getElementById("stats").innerHTML=
    `<div class="stat">Total&nbsp;: <strong>${totalCount}</strong></div>
     <div class="stat">Aperçu affiché&nbsp;: <strong>${allData.length}</strong></div>`;
  document.getElementById("subtitle").textContent="— "+totalCount+" opérations";
  document.getElementById("tbody").innerHTML=allData.map(r=>`<tr>
    <td>${r.departement||"—"}</td>
    <td><strong>${r.ville||"—"}</strong></td>
    <td>${r.operation||"—"}</td>
    <td><span class="badge-ref">${r.referentiel||"—"}</span></td>
    <td>${r.mentions||"—"}</td>
    <td>${r.fin_validite||"—"}</td>
  </tr>`).join("");
  document.getElementById("tfoot-msg").textContent =
    totalCount>20 ? "Affichage des 20 premières lignes sur "+totalCount+" — téléchargez l'Excel pour tout voir." : "";
  document.getElementById("results-card").style.display="block";
}
function telecharger(){ window.location.href="/api/download"; }
function setStatus(type,msg){
  const el=document.getElementById("status-box");
  el.className=type==="loading"?"s-loading":type==="done"?"s-done":"s-error";
  el.innerHTML=msg; el.style.display="block";
}
function setProgress(pct){
  document.getElementById("pw").style.display="block";
  document.getElementById("pb").style.width=pct+"%";
  if(pct>=100) setTimeout(()=>document.getElementById("pw").style.display="none",800);
}
function reset(){
  document.getElementById("f-op").value="";
  document.getElementById("f-ville").value="";
  ["ms-ref","ms-ment","ms-dept"].forEach(id=>{
    document.querySelectorAll('#'+id+' input:checked').forEach(c=>c.checked=false);
    msLabel(id);
  });
  document.getElementById("results-card").style.display="none";
  document.getElementById("status-box").style.display="none";
  document.getElementById("btn-dl").disabled=true;
  allData=[]; totalCount=0;
}
</script>
</body></html>'''

HTML = (HTML.replace("__REF__",  _checkboxes("ms-ref",  REFERENTIELS))
            .replace("__MENT__", _checkboxes("ms-ment", MENTIONS))
            .replace("__DEPT__", _checkboxes("ms-dept", DEPARTEMENTS)))


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5001))
    is_local = port == 5001
    print("=" * 50)
    print("  Extracteur Prestaterre BEE")
    print(f"  http://localhost:{port}")
    print("=" * 50)
    if is_local:
        import webbrowser
        threading.Timer(1.5, lambda: webbrowser.open(f"http://localhost:{port}")).start()
    app.run(host="0.0.0.0", port=port, debug=False)